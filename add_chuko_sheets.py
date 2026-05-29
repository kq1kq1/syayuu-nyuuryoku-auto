"""
既存の社有入力テンプレート.xlsx に中古用シートを追加するスクリプト。
既存シートのデータは一切変更しない。
すでに中古シートが存在する場合はスキップする。

追加シート:
  写真（中古）        : 外観5枚(01-05.jpg) + 内観21枚(06-26.jpg) = 26行
  売主コメント（中古）: 12行
  ホームズ_外観（中古）: 5行
  ホームズ_内観（中古）: 21行
  ピタクラ（中古）    : 26行
  スカイヤーズ（中古）: 26行
"""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# ===== カラー定数 =====
NAVY  = '1A5276'
WHITE = 'FFFFFF'
BLUE  = 'E3F2FD'
LIME  = 'F1F8E9'   # 中古用アクセント


def _hdr(cell, text):
    cell.value = text
    cell.font = Font(name='Arial', bold=True, color=WHITE, size=10)
    cell.fill = PatternFill('solid', fgColor=NAVY)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def _fill_sheet(ws, headers, rows, accent_color):
    """汎用シート作成（ヘッダー行 + データ行）"""
    for c, h in enumerate(headers, 1):
        _hdr(ws.cell(1, c), h)
    ws.row_dimensions[1].height = 22

    for r, vals in enumerate(rows, 2):
        fill = WHITE if r % 2 == 0 else accent_color
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c)
            cell.value = v
            cell.font = Font(name='Arial', size=10)
            cell.fill = PatternFill('solid', fgColor=fill)
            cell.alignment = Alignment(vertical='top', wrap_text=True)

    col_widths = [8, 24, 26, 42, 16]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w


# ===== 中古写真行定義 =====
_GAIKAN_ROWS = [(i, f'{i:02d}.jpg', '', '', '外観') for i in range(1, 6)]   # 01-05
_NAIKAN_ROWS = [(i, f'{i:02d}.jpg', '', '', '内観') for i in range(6, 27)]  # 06-26
_ALL_ROWS    = _GAIKAN_ROWS + _NAIKAN_ROWS  # 26行

_SUUMO_HEADERS  = ['順番', 'ファイル名', 'キャプション', '文言', '備考']
_HOMES_HEADERS  = ['順番', 'ファイル名', 'キャプション', 'コメント（文言）', '備考']

# 追加する中古シートの定義: (シート名, ヘッダー, 行データ, アクセントカラー)
CHUKO_SHEETS = [
    ('写真（中古）',         _SUUMO_HEADERS, _ALL_ROWS,    LIME),
    ('売主コメント（中古）', _SUUMO_HEADERS,
     [(i, f'{i:02d}.jpg', '', '', '') for i in range(1, 13)], LIME),
    ('ホームズ_外観（中古）', _HOMES_HEADERS, _GAIKAN_ROWS, LIME),
    ('ホームズ_内観（中古）', _HOMES_HEADERS, _NAIKAN_ROWS, LIME),
    ('ピタクラ（中古）',     _SUUMO_HEADERS, _ALL_ROWS,    LIME),
    ('スカイヤーズ（中古）', _SUUMO_HEADERS, _ALL_ROWS,    LIME),
]


def add_chuko_sheets(filepath: str) -> list[str]:
    """
    指定したExcelファイルに中古シートを追加する。
    戻り値: 追加したシート名のリスト（スキップしたものは含まない）
    """
    wb = openpyxl.load_workbook(filepath)
    existing = set(wb.sheetnames)
    added = []

    for sheet_name, headers, rows, color in CHUKO_SHEETS:
        if sheet_name in existing:
            print(f"  スキップ（既存）: {sheet_name}")
            continue
        ws = wb.create_sheet(sheet_name)
        _fill_sheet(ws, headers, rows, color)
        added.append(sheet_name)
        print(f"  追加: {sheet_name}（{len(rows)}行）")

    wb.save(filepath)
    return added


if __name__ == '__main__':
    path = sys.argv[1] if len(sys.argv) > 1 else '社有入力テンプレート.xlsx'
    print(f"対象ファイル: {path}")
    added = add_chuko_sheets(path)
    if added:
        print(f"完了: {len(added)}シートを追加しました")
    else:
        print("追加するシートはありませんでした（すべて既存）")
