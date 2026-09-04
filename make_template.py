import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

NAVY   = '1A5276'
YELLOW = 'FFFDE7'
GREEN  = 'E8F5E9'
BLUE   = 'E3F2FD'
WHITE  = 'FFFFFF'

wb = openpyxl.Workbook()

def hdr(cell, text):
    cell.value = text
    cell.font = Font(name='Arial', bold=True, color=WHITE, size=10)
    cell.fill = PatternFill('solid', fgColor=NAVY)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

def set_row(ws, row, label, val, fill, height=None):
    a = ws.cell(row, 1)
    b = ws.cell(row, 2)
    a.value = label
    a.font = Font(name='Arial', bold=True, size=10)
    a.fill = PatternFill('solid', fgColor=fill)
    a.alignment = Alignment(vertical='top', wrap_text=True)
    b.value = val
    b.font = Font(name='Arial', size=10)
    b.fill = PatternFill('solid', fgColor=fill)
    b.alignment = Alignment(vertical='top', wrap_text=True)
    if height:
        ws.row_dimensions[row].height = height

def sec_hdr(ws, row, text):
    ws.merge_cells(f'A{row}:C{row}')
    c = ws.cell(row, 1)
    c.value = text
    c.font = Font(name='Arial', bold=True, color=WHITE, size=11)
    c.fill = PatternFill('solid', fgColor=NAVY)
    c.alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[row].height = 24

# ===== Sheet1: 基本情報 =====
ws1 = wb.active
ws1.title = '基本情報'
hdr(ws1['A1'], 'フィールド名')
hdr(ws1['B1'], '値')
ws1.row_dimensions[1].height = 22

fields = [
    ('担当者名',       '神邑 尚志'),
    ('リンク表示名称1', '天空サイト'),
    ('リンク先URL1',   'https://www.i-unit.co.jp/'),
    ('リンク表示名称2', 'ピタットハウスHP'),
    ('リンク先URL2',   'https://www.i-unit.jp/'),
    ('リンク表示名称3', '天空の茶〜販売予告〜'),
    ('リンク先URL3',   'https://www.i-unit.co.jp/jst/notice/'),
    ('リンク表示名称4', 'スタッフブログ'),
    ('リンク先URL4',   'http://www.i-unitp.jp/blog/'),
    ('リンク表示名称5', '会社案内'),
    ('リンク先URL5',   'http://www.i-unit.jp/profile/'),
]
for r, (k, v) in enumerate(fields, 2):
    set_row(ws1, r, k, v, YELLOW)

ws1.column_dimensions['A'].width = 22
ws1.column_dimensions['B'].width = 48

# ===== Sheet2: 写真 =====
ws2 = wb.create_sheet('写真')
for c, h in enumerate(['順番','ファイル名','キャプション','文言','スーモスロット参考','キャプション（土地）'], 1):
    hdr(ws2.cell(1, c), h)
ws2.row_dimensions[1].height = 22

photos = [
    (1,'01_外観.jpg',      '現地外観写真',          '外観イメージ',          '先先①'),
    (2,'02_間取り.jpg',    '間取り図',              '建物面積88.93㎡',        '先先②'),
    (3,'03_区画図.jpg',    '区画図',                '土地面積128.73㎡',       '先先③'),
    (4,'04_外観2.jpg',     '現地外観写真',          '外観イメージ②',         '画像1'),
    (5,'05_リビング.jpg',  '同仕様写真（リビング）', '【同仕様】広々リビング', '画像2'),
    (6,'06_キッチン.jpg',  '同仕様写真（キッチン）', '【同仕様】キッチン',     '画像3'),
    (7,'07_洋室.jpg',      '同仕様写真（洋室）',    '【同仕様】洋室',         '画像4'),
    (8,'08_浴室.jpg',      '同仕様写真（浴室）',    '【同仕様】浴室',         '画像5'),
]
for i in range(9, 21):
    slot = f'画像{i-3}' if i <= 15 else ''
    photos.append((i, '', '', '', slot))

for r, vals in enumerate(photos, 2):
    fill = WHITE if r % 2 == 0 else BLUE
    for c, v in enumerate(vals, 1):
        cell = ws2.cell(r, c)
        cell.value = v
        cell.font = Font(name='Arial', size=10)
        cell.fill = PatternFill('solid', fgColor=fill)
        cell.alignment = Alignment(vertical='top', wrap_text=True)

ws2.column_dimensions['A'].width = 8
ws2.column_dimensions['B'].width = 24
ws2.column_dimensions['C'].width = 26
ws2.column_dimensions['D'].width = 42
ws2.column_dimensions['E'].width = 16
ws2.column_dimensions['F'].width = 26

# ===== Sheet3: 支払い例 =====
ws3 = wb.create_sheet('支払い例')

sec_hdr(ws3, 1, '▼ 入力項目（物件ごとに変更してください）')
set_row(ws3, 2, '物件価格（万円）',   4798,  YELLOW)
set_row(ws3, 3, '変動金利（年利 %）', 0.725, YELLOW)
set_row(ws3, 4, '返済期間（年）',     40,    YELLOW)

for row, note in [(2,'← 毎回変更'), (3,'← 金利変動時に変更'), (4,'← 通常40年固定')]:
    c = ws3.cell(row, 3)
    c.value = note
    c.font = Font(name='Arial', size=9, color='888888')
    c.alignment = Alignment(vertical='top')

sec_hdr(ws3, 6, '▼ 計算結果（自動計算）')
set_row(ws3, 7, '月々返済額（円）',
    '=ROUND(B2*10000*(B3/100/12)/(1-(1+B3/100/12)^(-B4*12)),0)', GREEN)
ws3.cell(7, 2).number_format = '#,##0'
ws3.cell(7, 3).value = '元利均等返済'
ws3.cell(7, 3).font = Font(name='Arial', size=9, color='888888')

sec_hdr(ws3, 9, '▼ 入力テキスト（ツールがそのままサイトに自動入力します）')

set_row(ws3, 10, '物件(住戸)情報\n(shriBknJoho)',
    '【物件金額100%のお借入の場合】', BLUE)

set_row(ws3, 11, '金額\n(shriKingaku)',
    '="価格"&B2&"万円、頭金0万円、借入額"&B2&"万円"', BLUE)

set_row(ws3, 12, '金利（示示）\n(shriKinri)',
    '="年利"&TEXT(B3,"0.000")&"% 変動金利 返済期間"&B4&"年"', BLUE)

set_row(ws3, 13, '支払額\n(shriGakubu)',
    '="毎月"&TEXT(B7,"#,##0")&"円 ボーナス時加算 0円 →団体信用生命保険付き"', BLUE)

loan_text = (
    '・紹介金融機関／弊社提携銀行(千葉銀行・常陽銀行・住信SBIネット銀行)\n'
    '・販売価格に対する融資限度額／100％以内\n'
    '・初期費用のお借入れも可\n'
    '・年齢制限/20歳以上65歳以下で完済時満80歳未満の方\n'
    '・返済期間／1年～40年\n'
    '・利率／年利0.725％～1.045％（変動金利）※弊社提携特別金利\n'
    '・団信付き・疾病特約\n'
    '・保証料／銀行指定による(お借入金額0.2％～0.4％)\n'
    '※適用される金利は、融資実行時のものとなり、表記されている金利と異なる場合があります。'
)
set_row(ws3, 14, '住宅ローンのご案内\n(shriLoan)', loan_text, BLUE, height=130)

flat_text = (
    '最長35年の長期固定金利住宅ローンです。\n'
    '保証料0円、繰上返済手数料0円\n'
    '・紹介金融機関／提携銀行他\n'
    '・販売価格に対する融資限度の割合／100％以内\n'
    '・年収に対する年間融資額の割合／30％～35％（諸費用も可）\n'
    '・返済期間／35年以内\n'
    '・利率／固定金利\n'
    '・保証料／なし\n'
    '・事務手数料／33000円・お借入金額の2.2％'
)
set_row(ws3, 15, 'フラット35ローンご案内\n(loanAnnai)', flat_text, BLUE, height=130)

ws3.column_dimensions['A'].width = 24
ws3.column_dimensions['B'].width = 52
ws3.column_dimensions['C'].width = 20

# ===== Sheet4: 売主コメント =====
ws4 = wb.create_sheet('売主コメント')
for c, h in enumerate(['順番','ファイル名','キャプション','文言','スーモスロット参考','キャプション（土地）'], 1):
    hdr(ws4.cell(1, c), h)
ws4.row_dimensions[1].height = 22

baishuu_photos = [
    (1, '01_外観.jpg',     '現地外観写真',           '外観イメージ',           '売主①'),
    (2, '02_間取り.jpg',   '間取り図',               '建物面積88.93㎡',         '売主②'),
    (3, '03_区画図.jpg',   '区画図',                 '土地面積128.73㎡',        '売主③'),
]
for i in range(4, 13):
    baishuu_photos.append((i, '', '', '', f'売主{i}'))

for r, vals in enumerate(baishuu_photos, 2):
    fill = WHITE if r % 2 == 0 else GREEN
    for c, v in enumerate(vals, 1):
        cell = ws4.cell(r, c)
        cell.value = v
        cell.font = Font(name='Arial', size=10)
        cell.fill = PatternFill('solid', fgColor=fill)
        cell.alignment = Alignment(vertical='top', wrap_text=True)

ws4.column_dimensions['A'].width = 8
ws4.column_dimensions['B'].width = 24
ws4.column_dimensions['C'].width = 26
ws4.column_dimensions['D'].width = 42
ws4.column_dimensions['E'].width = 16
ws4.column_dimensions['F'].width = 26

# ===== Sheet5: レイアウト指定 =====
ws5 = wb.create_sheet('レイアウト指定')
for c, h in enumerate(['段目', 'パターン', 'タイトル'], 1):
    hdr(ws5.cell(1, c), h)
ws5.row_dimensions[1].height = 22

layout_examples = [
    (1, 'A', ''),
    (2, 'A', ''),
    (3, 'A', ''),
]
for r, vals in enumerate(layout_examples, 2):
    fill = WHITE if r % 2 == 0 else YELLOW
    for c, v in enumerate(vals, 1):
        cell = ws5.cell(r, c)
        cell.value = v
        cell.font = Font(name='Arial', size=10)
        cell.fill = PatternFill('solid', fgColor=fill)
        cell.alignment = Alignment(vertical='top', wrap_text=True)

ws5.column_dimensions['A'].width = 8
ws5.column_dimensions['B'].width = 12
ws5.column_dimensions['C'].width = 48

def make_homes_photo_sheet(wb, sheet_name, example_rows, accent_color=None):
    """ホームズ用5列写真シートを作成（順番/ファイル名/キャプション/文言/備考）"""
    if accent_color is None:
        accent_color = BLUE
    ws = wb.create_sheet(sheet_name)
    for c, h in enumerate(['順番', 'ファイル名', 'キャプション', 'コメント（文言）', '備考'], 1):
        hdr(ws.cell(1, c), h)
    ws.row_dimensions[1].height = 22

    for r, vals in enumerate(example_rows, 2):
        fill = WHITE if r % 2 == 0 else accent_color
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c)
            cell.value = v
            cell.font = Font(name='Arial', size=10)
            cell.fill = PatternFill('solid', fgColor=fill)
            cell.alignment = Alignment(vertical='top', wrap_text=True)

    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 24
    ws.column_dimensions['C'].width = 26
    ws.column_dimensions['D'].width = 42
    ws.column_dimensions['E'].width = 16

# ===== Sheet6: ホームズ用（基本画像） =====
homes_photos = [(i, '', '', '', '') for i in range(1, 21)]
make_homes_photo_sheet(wb, 'ホームズ用', homes_photos)

# ===== Sheet7: ホームズ_外観分譲地 =====
gaigai_photos = [(i, '', '', '', '') for i in range(1, 11)]
make_homes_photo_sheet(wb, 'ホームズ_外観分譲地', gaigai_photos)

# ===== Sheet8: ホームズ_内観 =====
naikan_photos = [(i, '', '', '', '') for i in range(1, 21)]
make_homes_photo_sheet(wb, 'ホームズ_内観', naikan_photos)

# ===== Sheet9: ピタクラ =====
def make_site_photo_sheet(wb, sheet_name, rows, accent_color, tochi_column=False):
    """汎用5列写真シート（順番/ファイル名/キャプション/文言/備考）

    tochi_column=True で「キャプション（土地）」列を右端に足す。
    土地は新築と同じシートを使い、この列がある行だけキャプションを差し替える。
    """
    ws = wb.create_sheet(sheet_name)
    headers = ['順番', 'ファイル名', 'キャプション', '文言', '備考']
    if tochi_column:
        headers.append('キャプション（土地）')
    for c, h in enumerate(headers, 1):
        hdr(ws.cell(1, c), h)
    ws.row_dimensions[1].height = 22
    for r, vals in enumerate(rows, 2):
        fill = WHITE if r % 2 == 0 else accent_color
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c)
            cell.value = v
            cell.font = Font(name='Arial', size=10)
            cell.fill = PatternFill('solid', fgColor=fill)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 24
    ws.column_dimensions['C'].width = 26
    ws.column_dimensions['D'].width = 42
    ws.column_dimensions['E'].width = 16
    if tochi_column:
        ws.column_dimensions['F'].width = 26

ORANGE = 'FFF3E0'
PURPLE = 'F3E5F5'

pitakura_rows = [(i, '', '', '', '') for i in range(1, 21)]
make_site_photo_sheet(wb, 'ピタクラ', pitakura_rows, ORANGE, tochi_column=True)

# ===== Sheet10: スカイヤーズ =====
skyyers_rows = [(i, '', '', '', '') for i in range(1, 21)]
make_site_photo_sheet(wb, 'スカイヤーズ', skyyers_rows, PURPLE)

# ===== 中古物件用シート =====
# 外観5枚(01〜05) + 内観21枚(06〜26) = 計26枚
# ファイル名は連番リネーム後の名前（01.jpg〜26.jpg）と一致させる
LIME = 'F1F8E9'  # 中古用アクセントカラー（薄緑）

_chuko_gaikan = [(i, f'{i:02d}.jpg', '', '', '外観') for i in range(1, 6)]
_chuko_naikan = [(i, f'{i:02d}.jpg', '', '', '内観') for i in range(6, 27)]
_chuko_all    = _chuko_gaikan + _chuko_naikan

# ===== Sheet11: 写真（中古） =====
make_site_photo_sheet(wb, '写真（中古）', _chuko_all, LIME)

# ===== Sheet12: 売主コメント（中古） =====
_baishuu_chuko = [(i, f'{i:02d}.jpg', '', '', '') for i in range(1, 13)]
make_site_photo_sheet(wb, '売主コメント（中古）', _baishuu_chuko, LIME)

# ===== Sheet13: ホームズ_外観（中古）=====
make_homes_photo_sheet(wb, 'ホームズ_外観（中古）', _chuko_gaikan, LIME)

# ===== Sheet14: ホームズ_内観（中古）=====
make_homes_photo_sheet(wb, 'ホームズ_内観（中古）', _chuko_naikan, LIME)

# ===== Sheet15: ピタクラ（中古）=====
make_site_photo_sheet(wb, 'ピタクラ（中古）', _chuko_all, LIME)

# ===== Sheet16: スカイヤーズ（中古）=====
make_site_photo_sheet(wb, 'スカイヤーズ（中古）', _chuko_all, LIME)

# 保存はスクリプトとして実行されたときだけ行う。
# ガードが無いと、import しただけでカレントディレクトリに
# xlsx を書き出してしまい、既存テンプレートを壊す恐れがある。
if __name__ == '__main__':
    out = sys.argv[1] if len(sys.argv) > 1 else '社有入力テンプレート.xlsx'
    wb.save(out)
    print('OK:', out)
